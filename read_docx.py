
import sys
import os

try:
    from docx import Document
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

doc_path = r"docs\Vector_OS_Cross_Platform_Architecture_Design_v5.docx"

if not os.path.exists(doc_path):
    print(f"Document not found: {doc_path}")
    sys.exit(1)

print(f"Reading document: {doc_path}\n")
print("=" * 80)

doc = Document(doc_path)

for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)
        print()

print("=" * 80)
print("\nTables:")
for i, table in enumerate(doc.tables):
    print(f"\nTable {i+1}:")
    for row in table.rows:
        row_text = [cell.text.strip() for cell in row.cells]
        print(" | ".join(row_text))

